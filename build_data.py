import openpyxl, json, re, unicodedata, os
src='/mnt/data/files/Draft Oyuncu Listesi(1).xlsx'
wb=openpyxl.load_workbook(src,data_only=True)
ws=wb['130 Oyuncu']
catmap={
 'Kaleci':'GK','Sağ Bek':'RB','Sol Bek':'LB','Stoper':'CB','Ön Libero':'DM','Orta Saha':'CM','10 Numara':'CAM','Sağ Açık':'RW','Sol Açık':'LW','Santrfor':'ST'
}
players=[]
left_cat=None; right_cat=None
for r in range(1,ws.max_row+1):
    a,b,c,d,e,f,g,h,i,j,k,l,m=[ws.cell(r,col).value for col in range(1,14)]
    if isinstance(a,str) and a.strip() in catmap: left_cat=catmap[a.strip()]
    if isinstance(h,str) and h.strip() in catmap: right_cat=catmap[h.strip()]
    # left player row if rating numeric and category set
    if left_cat and isinstance(a,str) and isinstance(c,(int,float)) and isinstance(d,(int,float)):
        players.append({'name':a.strip(),'country':str(b).strip() if b else '', 'rating':int(c),'fee':int(d),'position':left_cat,'link':str(e).strip() if e else None,'linkBonus':int(f) if isinstance(f,(int,float)) else 0})
    if right_cat and isinstance(h,str) and isinstance(j,(int,float)) and isinstance(k,(int,float)):
        players.append({'name':h.strip(),'country':str(i).strip() if i else '', 'rating':int(j),'fee':int(k),'position':right_cat,'link':str(l).strip() if l else None,'linkBonus':int(m) if isinstance(m,(int,float)) else 0})
# de-duplicate by name+position if any, stable
seen=set(); out=[]
for p in players:
    key=(p['name'].casefold(),p['position'])
    if key in seen: continue
    seen.add(key)
    p['id']=f"p{len(out)+1:03d}"
    out.append(p)
assert len(out)==130, f'expected 130 players, got {len(out)}'
with open('/mnt/data/futbol-draft-yeni/players.json','w',encoding='utf-8') as f: json.dump({'source':'Draft Oyuncu Listesi(1).xlsx','players':out},f,ensure_ascii=False,indent=2)
from collections import Counter
print(len(out),Counter(p['position'] for p in out))
for p in out:
 if p['fee'] != (p['rating']-69)*5 or p['fee']>150: print('PRICE CHECK',p)
